"""Exercise downloaded credentials, including the original bytes for render QA.

DOCX_QA_PRIVATE_DIR is optional, private, and contains test credentials. Never
publish it. Only the encrypted payload and post-render redactions are public.
"""
import hashlib
import importlib
import io
import json
import os
from pathlib import Path

import pytest
from docx import Document
from CTFd.models import Users
from .fixtures import csrf, login

storage = importlib.import_module('CTFd.plugins.ctfd-account-provisioning.services.storage')
BASE = '/plugins/ctfd-account-provisioning/admin'


@pytest.mark.parametrize('lang', ['vi', 'en'])
def test_downloaded_cards_match_accounts_and_cached_retry(app, admin_client, lang):
    names = ['Nguyễn Văn An', 'Trần Thị Bình', 'Lê Quốc Việt', 'Nguyễn Văn An']
    csv = 'Player ID,Name\n' + ''.join(f'P{i:03d},{name}\n' for i, name in enumerate(names, 1))
    response = admin_client.post(BASE + '/upload', data={
        'file': (io.BytesIO(csv.encode('utf-8-sig')), 'participants.csv'),
        'email_domain': 'ctf.local', 'lang': lang,
        'nonce': csrf(admin_client)['CSRF-Token'],
    }, headers=csrf(admin_client))
    assert response.status_code == 200
    preview = response.get_json()
    assert preview['valid_count'] == 4 and preview['invalid_count'] == 0
    participants = preview['participants']
    usernames = [p['username'] for p in participants]
    assert len(set(usernames)) == 4
    with app.app_context():
        assert Users.query.filter(Users.name.in_(usernames)).count() == 0

    args = {'json': {'token': preview['token'], 'lang': lang}, 'headers': csrf(admin_client)}
    provision = admin_client.post(BASE + '/provision', **args)
    assert provision.status_code == 200
    result = provision.get_json()
    assert result['total_success'] == 4 and result['total_failed'] == 0
    retry = admin_client.post(BASE + '/provision', **args)
    assert retry.status_code == 200
    assert retry.get_json()['successful'] == result['successful']
    url = BASE + '/download/' + result['doc_token']
    first, second = admin_client.get(url), admin_client.get(url)
    assert first.status_code == second.status_code == 200
    assert first.data == second.data
    assert first.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'attachment' in first.headers['Content-Disposition']
    document = Document(io.BytesIO(first.data))
    assert len(document.tables) == 4
    assert [t.cell(1, 1).text for t in document.tables] == names
    assert [t.cell(2, 1).text for t in document.tables] == usernames
    assert all(t.cell(3, 1).text == p['password'] for t, p in zip(document.tables, participants))
    labels = ('Tên Đăng Nhập', 'Mật Khẩu') if lang == 'vi' else ('Username', 'Password')
    assert all((t.cell(2, 0).text, t.cell(3, 0).text) == labels for t in document.tables)
    for participant in participants:
        client = app.test_client()
        try:
            assert login(client, participant['username'], participant['password']).status_code == 302
            assert client.get('/settings').status_code == 200
        finally:
            client.get('/logout')
    with app.app_context():
        assert Users.query.filter(Users.name.in_(usernames)).count() == 4
        job = storage.ProvisioningJob.query.filter_by(token_hash=storage.digest(preview['token'])).one()
        encrypted_payload = job.payload

    target = os.environ.get('DOCX_QA_PRIVATE_DIR')
    if target:
        directory = Path(target)
        directory.mkdir(mode=0o700, parents=True, exist_ok=True)
        path = directory / f'download-{lang}.docx'
        path.write_bytes(first.data)
        path.chmod(0o600)
        (directory / f'download-{lang}.encrypted').write_bytes(encrypted_payload)
        (directory / f'download-{lang}.json').write_text(json.dumps({
            'language': lang, 'valid_rows': 4, 'accounts_created': 4,
            'successful_logins': 4, 'credential_cards': 4,
            'repeat_confirmation_same_users': True, 'repeat_download_same_bytes': True,
            'docx_sha256': hashlib.sha256(first.data).hexdigest(),
            'mime_type': first.mimetype, 'attachment_header': True,
            'original_bytes_preserved': True,
        }, indent=2) + '\n')
