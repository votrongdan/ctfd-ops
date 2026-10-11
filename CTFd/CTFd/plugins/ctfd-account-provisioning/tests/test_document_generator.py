"""
Tests for the credential document generator module.
"""

import os
import sys
from dataclasses import dataclass
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn

_plugin_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _plugin_dir not in sys.path:
    sys.path.insert(0, _plugin_dir)

from services.document_generator import (
    generate_credential_document,
)


@dataclass
class MockAccount:
    index: int = 0
    name: str = "Test User"
    username: str = "test_user"
    email: str = "test@ctf.local"
    password: str = "Pass123"
    user_id: int = 1


class TestGenerateCredentialDocument:

    def test_document_created(self):
        accounts = [MockAccount()]
        output = generate_credential_document(accounts, lang="en")
        assert isinstance(output, BytesIO)
        doc = Document(output)
        assert len(doc.paragraphs) > 0

    def test_correct_card_count(self):
        accounts = [MockAccount(name=f"User{i}", username=f"user{i}") for i in range(5)]
        output = generate_credential_document(accounts, lang="en")
        doc = Document(output)
        # Each card is a table with 5 rows
        assert len(doc.tables) == 5

    def test_vietnamese_name_preserved(self):
        accounts = [MockAccount(name="Nguyễn Văn An", username="nguyen_van_an")]
        output = generate_credential_document(accounts, lang="vi")
        doc = Document(output)
        all_text = "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        assert "Nguyễn Văn An" in all_text

    def test_english_language(self):
        accounts = [MockAccount()]
        output = generate_credential_document(accounts, lang="en")
        doc = Document(output)
        all_text = "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        assert "Username" in all_text
        assert "Password" in all_text

    def test_vietnamese_language(self):
        accounts = [MockAccount()]
        output = generate_credential_document(accounts, lang="vi")
        doc = Document(output)
        all_text = "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        assert "Tên Đăng Nhập" in all_text
        assert "Mật Khẩu" in all_text

    def test_correct_username_in_document(self):
        accounts = [MockAccount(username="nguyen_van_an")]
        output = generate_credential_document(accounts, lang="en")
        doc = Document(output)
        all_text = "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        assert "nguyen_van_an" in all_text

    def test_correct_password_in_document(self):
        accounts = [MockAccount(password="SecretPass99")]
        output = generate_credential_document(accounts, lang="en")
        doc = Document(output)
        all_text = "\n".join(p.text for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
        assert "SecretPass99" in all_text

    def test_empty_accounts_list(self):
        output = generate_credential_document([], lang="en")
        doc = Document(output)
        assert len(doc.tables) == 0

    def test_only_successful_accounts(self):
        """Verify we only get cards for accounts we pass in."""
        accounts = [
            MockAccount(name="Success1", username="s1"),
            MockAccount(name="Success2", username="s2"),
        ]
        output = generate_credential_document(accounts, lang="en")
        doc = Document(output)
        assert len(doc.tables) == 2

    @pytest.mark.parametrize("lang", ["en", "vi"])
    def test_cards_keep_rows_together_but_allow_breaks_between_cards(self, lang):
        accounts = [MockAccount(name=f"Player {i}", username=f"player_{i}") for i in range(4)]
        doc = Document(generate_credential_document(accounts, lang=lang))

        for table in doc.tables:
            for index, row in enumerate(table.rows):
                assert row._tr.trPr.find(qn("w:cantSplit")) is not None
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        assert paragraph.paragraph_format.keep_together is True
                        assert paragraph.paragraph_format.keep_with_next is (index < 4)
        # Spacing should travel with the next card, without joining two cards.
        assert all(p.paragraph_format.keep_with_next for p in doc.paragraphs[1:])
