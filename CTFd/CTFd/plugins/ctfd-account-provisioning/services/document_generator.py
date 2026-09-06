"""
Credential document generator for CTFd Account Provisioning plugin.
Generates a Word (.docx) document with printable credential cards
for successfully created player accounts.
"""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    from ..i18n import get_text
except ImportError:
    from i18n import get_text


def _set_cell_border(cell, **kwargs):
    """
    Set cell borders. Usage:
        _set_cell_border(cell, top={"sz": 12, "val": "single", "color": "000000"}, ...)
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    for edge_name, edge_data in kwargs.items():
        element = OxmlElement(f"w:{edge_name}")
        for attr_name, attr_value in edge_data.items():
            element.set(qn(f"w:{attr_name}"), str(attr_value))
        tcPr.append(element)


def _set_cell_shading(cell, color):
    """Set background color of a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")
    tcPr.append(shading)


def _add_credential_card(doc, account, lang, card_index):
    """
    Add a single credential card to the document.

    :param doc: Document object
    :param account: ImportSuccess object
    :param lang: Language code
    :param card_index: Card number (for spacing)
    """
    # Add spacing between cards (except for the first one)
    if card_index > 0:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(6)
        spacer.paragraph_format.space_after = Pt(6)

    # Create a table for the credential card (acts as a bordered card)
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Set table width
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8)

    # --- Row 0: Header with card title ---
    header_cell = table.cell(0, 0)
    header_cell.merge(table.cell(0, 1))
    _set_cell_shading(header_cell, "2C3E50")
    header_para = header_cell.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_para.add_run(get_text("doc_title", lang))
    header_run.bold = True
    header_run.font.size = Pt(14)
    header_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    header_para.paragraph_format.space_before = Pt(8)
    header_para.paragraph_format.space_after = Pt(8)

    # --- Row 1: Player name ---
    name_label_cell = table.cell(1, 0)
    name_label_para = name_label_cell.paragraphs[0]
    name_label_run = name_label_para.add_run(get_text("card_name_label", lang))
    name_label_run.bold = True
    name_label_run.font.size = Pt(11)
    name_label_run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    name_value_cell = table.cell(1, 1)
    name_value_para = name_value_cell.paragraphs[0]
    name_run = name_value_para.add_run(account.name)
    name_run.bold = True
    name_run.font.size = Pt(13)

    # --- Row 2: Username ---
    user_label_cell = table.cell(2, 0)
    user_label_para = user_label_cell.paragraphs[0]
    user_label_run = user_label_para.add_run(get_text("card_username_label", lang))
    user_label_run.bold = True
    user_label_run.font.size = Pt(11)
    user_label_run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    user_value_cell = table.cell(2, 1)
    user_value_para = user_value_cell.paragraphs[0]
    user_run = user_value_para.add_run(account.username)
    user_run.bold = True
    user_run.font.size = Pt(14)
    user_run.font.name = "Consolas"
    user_run.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)

    # --- Row 3: Password ---
    pass_label_cell = table.cell(3, 0)
    pass_label_para = pass_label_cell.paragraphs[0]
    pass_label_run = pass_label_para.add_run(get_text("card_password_label", lang))
    pass_label_run.bold = True
    pass_label_run.font.size = Pt(11)
    pass_label_run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

    pass_value_cell = table.cell(3, 1)
    pass_value_para = pass_value_cell.paragraphs[0]
    pass_run = pass_value_para.add_run(account.password)
    pass_run.bold = True
    pass_run.font.size = Pt(14)
    pass_run.font.name = "Consolas"
    pass_run.font.color.rgb = RGBColor(0xE7, 0x4C, 0x3C)

    # --- Row 4: Note ---
    note_cell = table.cell(4, 0)
    note_cell.merge(table.cell(4, 1))
    _set_cell_shading(note_cell, "F8F9FA")
    note_para = note_cell.paragraphs[0]
    note_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note_para.add_run(get_text("card_note", lang))
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(0x95, 0xA5, 0xA6)

    # Apply borders to all cells
    border_settings = {
        "top": {"sz": "8", "val": "single", "color": "BDC3C7"},
        "bottom": {"sz": "8", "val": "single", "color": "BDC3C7"},
        "start": {"sz": "8", "val": "single", "color": "BDC3C7"},
        "end": {"sz": "8", "val": "single", "color": "BDC3C7"},
    }
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, **border_settings)


def generate_credential_document(successful_accounts, lang="en"):
    """
    Generate a Word document containing credential cards for
    successfully created accounts.

    :param successful_accounts: List of ImportSuccess objects
    :param lang: Language code ('en' or 'vi')
    :return: BytesIO object containing the .docx file
    """
    doc = Document()

    # Set default font to Calibri (supports Vietnamese Unicode)
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Document title
    title = doc.add_heading(get_text("doc_title", lang), level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Add each credential card
    for i, account in enumerate(successful_accounts):
        _add_credential_card(doc, account, lang, i)

    # Save to BytesIO
    output = BytesIO()
    doc.save(output)
    output.seek(0)

    return output
